from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = 'docs/DuPay_Quality_and_Security_Audit_2026-10-24.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.72)
sec.bottom_margin = Inches(.72)
sec.left_margin = Inches(.82)
sec.right_margin = Inches(.82)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10)
styles['Normal'].paragraph_format.space_after = Pt(7)
for name, size in [('Title', 20), ('Heading 1', 13), ('Heading 2', 11)]:
    s = styles[name]
    s.font.name = 'Aptos'
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0, 0, 0)

doc.add_paragraph('DuPay Quality and Security Audit', 'Title')
doc.add_paragraph('24 October 2026  |  Flutter mobile client  |  Code review and remediation record')
doc.add_paragraph('The review identified direct disclosure of credentials and transaction PINs through logs, sensitive account data in unencrypted preferences, a false authenticated state after password reset, and unsafe release signing fallback. Client-side fixes are implemented. This is a static review; a production or staging penetration test has not been performed.')

doc.add_heading('Scope and method', 1)
doc.add_paragraph('Reviewed Flutter source, Android and iOS configuration, request and response handling, local storage, authentication, money-moving flows, and existing tests. Searched for sensitive logging and cleartext transport settings. No backend source, staging credentials, emulator, or Flutter/Dart SDK was available. No traffic was sent to the production API.')

doc.add_heading('Findings and remediation', 1)
rows = [
 ('F1', 'Critical', 'Passwords, KYC data and transaction PINs printed by repositories', 'Removed request payload prints from authentication and payment repositories. Review remaining diagnostic output before release.'),
 ('F2', 'High', 'Network logger emitted payloads and errors with incomplete redaction', 'Logger now emits only method/path or status in debug mode; no headers, bodies, query strings or error details.'),
 ('F3', 'High', 'Profile, transactions, ePINs and foreign account records persisted in SharedPreferences', 'Sensitive cache keys now use process memory and remove legacy preference entries when accessed. Logout clears the memory cache.'),
 ('F4', 'High', 'Password reset stored user ID as access token and marked session authenticated', 'Reset clears stored credentials and provider session; the user must sign in again.'),
 ('F5', 'Medium', 'Release build could fall back to debug signing', 'Release task fails when signing properties are absent; no debug-signing fallback.'),
 ('F6', 'Medium', 'Unused transaction PIN storage methods allowed persistent raw PIN storage', 'Removed the methods. The client user model now retains only a PIN-set indicator.'),
]
table = doc.add_table(rows=1, cols=4)
table.style = 'Table Grid'
for c, v in zip(table.rows[0].cells, ['ID', 'Risk', 'Finding', 'Action']): c.text = v
for item in rows:
    for c, v in zip(table.add_row().cells, item): c.text = v
for row in table.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs: run.font.size = Pt(8)

doc.add_heading('Outstanding risks and validation', 1)
doc.add_paragraph('The backend appears to return a transaction_pin field in user responses. The client now discards its value, but the server should return only a boolean status. Existing sensitive preference data is deleted when each affected cache key is accessed or when cache clearing runs; users who do not launch the updated app may retain older local data until app data is cleared. Previously captured device logs cannot be recalled.')
doc.add_paragraph('Flutter analyze, Flutter tests, dependency audit, Android/iOS release builds, device inspection, and authenticated staging API penetration testing remain unverified because the required SDK and staging environment were unavailable. The current tests cover phone validation and selected widgets, but not authorization, account switching, logging, or payment flows. A full penetration test requires an isolated staging backend, test accounts, and explicit test boundaries.')

doc.add_heading('Recommended verification sequence', 1)
for text in [
 'Install Flutter/Dart and run flutter analyze and flutter test; resolve any failures.',
 'Build signed Android and iOS release artifacts and confirm the signing identity and backup behavior.',
 'On a test device, confirm login, registration, password reset, logout, account switching, and payment flows; inspect logs and preference stores for secrets.',
 'Against staging, test authorization boundaries, replay and idempotency, rate limiting, token expiry, PIN handling, and KYC upload validation.',
 'Change the backend user response to expose only has_transaction_pin, then remove the legacy client field.',
]: doc.add_paragraph(text, style='List Bullet')

doc.add_heading('Source locations', 1)
doc.add_paragraph('lib/repositories/auth_repository.dart; lib/repositories/transfer_repository.dart; lib/core/interceptors/logging_interceptor.dart; lib/services/cache/cache_service.dart; lib/services/storage/secure_storage_service.dart; lib/models/auth/user_model.dart; lib/providers/auth_provider.dart; android/app/build.gradle.kts.')
doc.save(OUT)
print(OUT)
